from io import BytesIO
from decimal import Decimal

from django import forms
from django.contrib import admin, messages
from django.contrib.admin.models import LogEntry
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import get_object_or_404
from django.template.response import TemplateResponse
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import ControllerReadingSubmission, Meter, Reading


REVIEW_THRESHOLD_M3 = Decimal('100')


def _import_meter_id(meter):
    marker = 'ID импорта:'
    for line in (meter.notes or '').splitlines():
        if marker in line:
            return line.split(marker, 1)[1].strip().split(';', 1)[0].strip()
    return ''


def _xlsx_safe(value):
    if value is None:
        return ''
    text = str(value)
    if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + text
    return value


def _reading_source(reading):
    submissions = list(reading.controller_submissions.all())
    if submissions:
        parts = []
        for item in submissions:
            actor = item.submitted_by.get_full_name() or item.submitted_by.username
            parts.append(f'Контролёр: {actor}; статус: {item.get_status_display()}')
        return ' | '.join(parts)
    if 'техническая дата по согласованной схеме' in (reading.notes or '').casefold():
        return 'Исторический импорт с технической датой'
    return 'Ручная / административная запись'


def export_readings_xlsx(request):
    if not request.user.has_perm('water.export_reading'):
        raise PermissionDenied

    queryset = Reading.objects.select_related(
        'meter', 'meter__account', 'meter__group', 'meter__node',
    ).prefetch_related(
        'controller_submissions__submitted_by',
    ).order_by('meter_id', 'date', 'id')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Показания'
    headers = [
        'Reading ID', 'Import / meter ID', 'Счётчик', 'Назначение',
        'Лицевой счёт', 'Участок / адрес', 'Группа', 'Узел',
        'Дата', 'Показание, м³', 'Расход от предыдущего, м³',
        'Контроль', 'Источник', 'Примечание',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:N1'

    previous_by_meter = {}
    count = 0
    for reading in queryset:
        previous = previous_by_meter.get(reading.meter_id)
        consumption = reading.value - previous.value if previous is not None else None
        previous_by_meter[reading.meter_id] = reading

        review = ''
        if reading.meter.kind == 'individual' and consumption is not None and consumption > REVIEW_THRESHOLD_M3:
            review = f'ПРОВЕРИТЬ: расход > {REVIEW_THRESHOLD_M3} м³'

        account = reading.meter.account
        row = [
            reading.pk,
            _import_meter_id(reading.meter),
            reading.meter.serial,
            reading.meter.get_kind_display(),
            account.number if account else '',
            account.plot if account else '',
            reading.meter.group.name if reading.meter.group else '',
            reading.meter.node.name,
            reading.date,
            reading.value,
            consumption if consumption is not None else '',
            review,
            _reading_source(reading),
            reading.notes,
        ]
        ws.append([_xlsx_safe(value) for value in row])
        count += 1

    for row in ws.iter_rows(min_row=2):
        row[8].number_format = 'dd.mm.yyyy'
        row[9].number_format = '0.000'
        if isinstance(row[10].value, (int, float, Decimal)):
            row[10].number_format = '0.000'

    widths = [12, 26, 24, 18, 20, 28, 24, 24, 14, 18, 24, 34, 38, 60]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    info = wb.create_sheet('Пояснение')
    info.append(['Назначение', 'Выгрузка всей таблицы показаний для визуального контроля.'])
    info.append(['Флаг ПРОВЕРИТЬ', f'Только для индивидуальных счётчиков: расход между соседними показаниями больше {REVIEW_THRESHOLD_M3} м³. Это сигнал для ручной проверки, а не утверждение об ошибке.'])
    info.append(['Исправления', 'Исправлять привязку счётчика следует через отдельное админское действие с подтверждением и причиной.'])
    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 110

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    LogEntry.objects.create(
        user=request.user,
        content_type=ContentType.objects.get_for_model(Reading),
        object_id='',
        object_repr='Выгрузка всей таблицы показаний XLSX',
        action_flag=2,
        change_message=f'Экспорт XLSX: {count} записей',
    )

    response = HttpResponse(
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="trud-water-readings.xlsx"'
    response['Cache-Control'] = 'no-store'
    return response


class MeterChoiceField(forms.ModelChoiceField):
    def label_from_instance(self, meter):
        account = meter.account
        place = account.plot if account else meter.group or meter.node
        import_id = _import_meter_id(meter)
        prefix = f'{import_id} · ' if import_id else ''
        return f'{prefix}{meter.serial} · {place}'


class ReassignReadingForm(forms.Form):
    destination_meter = MeterChoiceField(label='Правильный счётчик', queryset=Meter.objects.none())
    reason = forms.CharField(
        label='Причина исправления', min_length=3,
        widget=forms.Textarea(attrs={'rows': 3}),
        help_text='Причина попадёт в историю изменений и журнал администратора.',
    )
    reading_version = forms.IntegerField(widget=forms.HiddenInput())

    def __init__(self, *args, reading, **kwargs):
        super().__init__(*args, **kwargs)
        self.reading = reading
        self.fields['reading_version'].initial = reading.version
        self.fields['destination_meter'].queryset = Meter.objects.select_related(
            'account', 'group', 'node',
        ).filter(kind=reading.meter.kind).exclude(pk=reading.meter_id).order_by(
            'account__plot', 'group__name', 'serial', 'id',
        )

    def clean(self):
        data = super().clean()
        if data.get('reading_version') != self.reading.version:
            raise forms.ValidationError('Показание уже изменилось. Обновите страницу и повторите проверку.')
        return data


def validate_reassignment(reading, destination_meter):
    if destination_meter.pk == reading.meter_id:
        raise ValidationError('Выбран тот же счётчик.')
    if destination_meter.kind != reading.meter.kind:
        raise ValidationError('Назначение исходного и нового счётчика должно совпадать.')

    candidate = Reading(
        meter=destination_meter,
        date=reading.date,
        value=reading.value,
        notes=reading.notes,
    )
    candidate.full_clean()

    previous = Reading.objects.filter(
        meter=destination_meter, date__lt=reading.date,
    ).order_by('-date', '-id').first()
    following = Reading.objects.filter(
        meter=destination_meter, date__gt=reading.date,
    ).order_by('date', 'id').first()
    return previous, following


def reassign_reading(*, reading_id, destination_meter_id, reason, actor, expected_version):
    reason = (reason or '').strip()
    if not reason:
        raise ValidationError('Укажите причину исправления.')
    if not actor or not actor.is_superuser:
        raise PermissionDenied

    with transaction.atomic():
        reading = Reading.objects.select_for_update().select_related('meter').get(pk=reading_id)
        if reading.version != expected_version:
            raise ValidationError('Показание уже изменилось. Обновите страницу и повторите проверку.')
        destination = Meter.objects.select_for_update().get(pk=destination_meter_id)
        source = Meter.objects.select_for_update().get(pk=reading.meter_id)
        validate_reassignment(reading, destination)

        old_id = reading.pk
        old_meter_label = str(source)
        destination_label = str(destination)
        full_audit_reason = f'Исправление привязки: {old_meter_label} → {destination_label}. {reason}'
        history_reason = (
            f'Исправление привязки Reading #{old_id}: meter {source.pk}→{destination.pk}. {reason}'
        )[:100]

        replacement = Reading(
            meter=destination,
            date=reading.date,
            value=reading.value,
            notes=reading.notes,
        )
        replacement._history_user = actor
        replacement._change_reason = history_reason
        replacement.save()

        submissions = list(
            ControllerReadingSubmission.objects.select_for_update().filter(reading=reading)
        )
        for submission in submissions:
            submission.meter = destination
            submission.reading = replacement
            submission._history_user = actor
            submission._change_reason = history_reason
            submission.save(update_fields=['meter', 'reading'])

        reading._history_user = actor
        reading._change_reason = history_reason
        reading.delete()

        return replacement, old_id, len(submissions), old_meter_label, destination_label, full_audit_reason


def reassign_reading_view(request, reading_id):
    if not request.user.is_superuser:
        raise PermissionDenied

    reading = get_object_or_404(
        Reading.objects.select_related('meter', 'meter__account', 'meter__group', 'meter__node'),
        pk=reading_id,
    )
    form = ReassignReadingForm(request.POST or None, reading=reading)
    preview = None

    if request.method == 'POST' and form.is_valid():
        destination = form.cleaned_data['destination_meter']
        try:
            previous, following = validate_reassignment(reading, destination)
        except ValidationError as error:
            form.add_error('destination_meter', '; '.join(error.messages))
        else:
            preview = {
                'destination': destination,
                'previous': previous,
                'following': following,
            }
            if request.POST.get('confirm') == 'yes':
                try:
                    replacement, old_id, linked_count, old_label, new_label, full_reason = reassign_reading(
                        reading_id=reading.pk,
                        destination_meter_id=destination.pk,
                        reason=form.cleaned_data['reason'],
                        actor=request.user,
                        expected_version=form.cleaned_data['reading_version'],
                    )
                except ValidationError as error:
                    form.add_error(None, '; '.join(error.messages))
                else:
                    LogEntry.objects.create(
                        user=request.user,
                        content_type=ContentType.objects.get_for_model(Reading),
                        object_id=str(replacement.pk),
                        object_repr=str(replacement),
                        action_flag=2,
                        change_message=(
                            f'Исправлена привязка показания: Reading #{old_id}; '
                            f'{old_label} → {new_label}; связанных заявок контролёра: {linked_count}; '
                            f'{full_reason}'
                        ),
                    )
                    messages.success(
                        request,
                        f'Привязка исправлена. Новая запись #{replacement.pk}; '
                        f'связанных заявок контролёра перенесено: {linked_count}.',
                    )
                    return HttpResponseRedirect(reverse('admin:water_reading_change', args=[replacement.pk]))

    context = {
        **admin.site.each_context(request),
        'title': f'Исправление привязки показания #{reading.pk}',
        'opts': Reading._meta,
        'reading': reading,
        'form': form,
        'preview': preview,
        'linked_submissions': reading.controller_submissions.select_related('submitted_by').all(),
    }
    return TemplateResponse(request, 'admin/water/reading/reassign.html', context)
