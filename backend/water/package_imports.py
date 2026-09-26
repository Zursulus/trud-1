"""Validation for the multi-sheet water import workbook.

This module deliberately performs no writes to the working registry. It is the
first gate for the staged import package prepared from the legacy water file.
"""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import io
import re
import zipfile

from django.core.exceptions import ValidationError

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_ROWS_PER_SHEET = 5000

SHEETS = {
    'Участки': ('plot_id', 'label', 'address', 'cadastral_number', 'contact_name', 'phone', 'source_sheet', 'source_row', 'status'),
    'Люди': ('person_id', 'full_name', 'phone', 'plot_id', 'role', 'notes'),
    'Узлы': ('node_name', 'notes'),
    'Группы': ('group_name', 'node_name', 'source', 'notes'),
    'Состав групп': ('plot_id', 'group_name', 'starts', 'ends', 'notes'),
    'Индивидуальные счетчики': ('meter_id', 'kind', 'node_name', 'group_name', 'plot_id', 'serial', 'notes'),
    'Общие и контрольные': ('meter_id', 'kind', 'node_name', 'group_name', 'serial', 'source_text', 'source_sheet', 'source_row', 'status'),
    'Показания': ('meter_id', 'date', 'value_m3', 'notes', 'source_sheet', 'source_row'),
}

TECHNICAL_DATE_MARKER = 'техническая дата по согласованной схеме 5/20'
TECHNICAL_MONTHS = {'июнь': 6, 'июль': 7, 'август': 8}


def _looks_like_phone_number(value):
    """Reject contact numbers accidentally mapped into the reading column."""
    compact = re.sub(r'[\s()+\-.]', '', _text(value))
    return compact.isdigit() and 10 <= len(compact) <= 12


def _note_says_exact_date_unknown(value):
    """Detect source notes that explicitly say the exact reading date is unknown."""
    note = _text(value).casefold()
    return 'дата неизвестна' in note or ('точн' in note and 'дат' in note and 'неизвест' in note)


def _parse_date_text(value):
    """Parse an XLSX-normalized date value for validation only."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    for parser in (
        lambda s: date.fromisoformat(s[:10]),
        lambda s: datetime.strptime(s, '%d.%m.%Y').date(),
    ):
        try:
            return parser(text)
        except (ValueError, TypeError):
            pass
    return None


def _technical_date_matches_note(value, note_value):
    """Allow only the explicitly agreed synthetic 5th/20th date scheme."""
    note = _text(note_value).casefold()
    if TECHNICAL_DATE_MARKER not in note:
        return False
    reading_date = _parse_date_text(value)
    if reading_date is None:
        return False

    month = next((number for name, number in TECHNICAL_MONTHS.items() if name in note), None)
    if month is None or reading_date.month != month:
        return False

    if 'предыдущее' in note or 'начальное' in note:
        expected_day = 5
    elif 'текущее' in note:
        expected_day = 20
    else:
        return False
    return reading_date.day == expected_day


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(value) for value in next(values))
    except StopIteration:
        return (), []
    rows = []
    for number, values in enumerate(values, 2):
        if number > MAX_ROWS_PER_SHEET + 1:
            raise ValidationError(f'{sheet.title}: больше {MAX_ROWS_PER_SHEET} строк.')
        if not any(value not in (None, '') for value in values):
            continue
        normalized = tuple(_text(value) for value in values)
        # Prepared workbooks may repeat a service header inside a table.
        # Recognize it by the complete column signature rather than by position.
        if normalized[:len(header)] == header:
            continue
        rows.append((number, dict(zip(header, normalized))))
    return header, rows


def inspect_water_package(upload):
    """Return a dry-run report for the prepared workbook; never writes models."""
    data = upload.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE:
        raise ValidationError('Файл больше 5 МБ.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_SIZE:
                raise ValidationError('Распакованный Excel слишком велик.')
    except zipfile.BadZipFile as error:
        raise ValidationError('Файл не является корректным XLSX.') from error

    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as error:
        raise ValidationError('Не удалось безопасно прочитать XLSX.') from error

    missing = [name for name in SHEETS if name not in workbook.sheetnames]
    if missing:
        workbook.close()
        raise ValidationError('Нет обязательных листов: ' + ', '.join(missing))

    parsed = {}
    blocking_issues = []
    review_notes = []
    counts = {}
    for name, required in SHEETS.items():
        header, rows = _rows(workbook[name])
        absent = [column for column in required if column not in header]
        if absent:
            blocking_issues.append(f'{name}: нет колонок {", ".join(absent)}.')
        parsed[name] = rows
        counts[name] = len(rows)
    workbook.close()

    def values(sheet, field):
        return [row.get(field, '') for _, row in parsed[sheet] if row.get(field, '')]

    def unique(sheet, field):
        found = values(sheet, field)
        for value, count in Counter(found).items():
            if count > 1:
                blocking_issues.append(f'{sheet}: {field} «{value}» повторяется {count} раза.')
        return set(found)

    plot_ids = unique('Участки', 'plot_id')
    unique('Участки', 'address')
    person_ids = unique('Люди', 'person_id')
    node_names = unique('Узлы', 'node_name')
    group_names = unique('Группы', 'group_name')
    individual_meter_ids = unique('Индивидуальные счетчики', 'meter_id')
    system_meter_ids = unique('Общие и контрольные', 'meter_id')
    meter_ids = individual_meter_ids | system_meter_ids
    if individual_meter_ids & system_meter_ids:
        blocking_issues.append('Одинаковый meter_id встречается в индивидуальных и общих счётчиках.')

    for number, row in parsed['Люди']:
        if row.get('plot_id') and row['plot_id'] not in plot_ids:
            blocking_issues.append(f'Люди строка {number}: неизвестный plot_id {row["plot_id"]}.')
    for number, row in parsed['Группы']:
        if row.get('node_name') not in node_names:
            blocking_issues.append(f'Группы строка {number}: неизвестный узел {row.get("node_name") or "(пусто)"}.')
    for number, row in parsed['Состав групп']:
        if row.get('plot_id') not in plot_ids:
            blocking_issues.append(f'Состав групп строка {number}: неизвестный plot_id {row.get("plot_id") or "(пусто)"}.')
        if row.get('group_name') not in group_names:
            blocking_issues.append(f'Состав групп строка {number}: неизвестная группа {row.get("group_name") or "(пусто)"}.')

    allowed_kinds = {'main', 'line', 'individual', 'irrigation'}
    for sheet in ('Индивидуальные счетчики', 'Общие и контрольные'):
        for number, row in parsed[sheet]:
            kind = row.get('kind')
            if kind not in allowed_kinds:
                blocking_issues.append(f'{sheet} строка {number}: неизвестный kind {kind or "(пусто)"}.')
            if row.get('node_name') not in node_names:
                blocking_issues.append(f'{sheet} строка {number}: неизвестный узел {row.get("node_name") or "(пусто)"}.')
            if kind == 'individual' and row.get('plot_id') not in plot_ids:
                blocking_issues.append(f'{sheet} строка {number}: неизвестный plot_id {row.get("plot_id") or "(пусто)"}.')
            if kind == 'line' and row.get('group_name') not in group_names:
                blocking_issues.append(f'{sheet} строка {number}: для линейного счётчика нужна существующая группа.')
            if kind in {'main', 'irrigation'} and row.get('group_name'):
                blocking_issues.append(f'{sheet} строка {number}: общий/поливочный счётчик не должен иметь группу.')

    reading_keys = set()
    dated_readings = 0
    technical_dated_readings = 0
    undated_values = 0
    empty_readings = 0
    for number, row in parsed['Показания']:
        meter_id = row.get('meter_id')
        if meter_id not in meter_ids:
            blocking_issues.append(f'Показания строка {number}: неизвестный meter_id {meter_id or "(пусто)"}.')

        raw_value = row.get('value_m3', '')
        value = None
        if raw_value:
            try:
                value = Decimal(raw_value.replace(',', '.'))
                if value < 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                blocking_issues.append(f'Показания строка {number}: некорректное значение {raw_value}.')
            else:
                if _looks_like_phone_number(raw_value):
                    blocking_issues.append(
                        f'Показания строка {number}: значение похоже на телефон, '
                        'а не на показание счётчика.'
                    )
                    value = None
        else:
            empty_readings += 1
            review_notes.append(
                f'Показания строка {number}: значение отсутствует в источнике; показание не создаётся.'
            )

        reading_date = row.get('date')
        if reading_date and _note_says_exact_date_unknown(row.get('notes')):
            if not _technical_date_matches_note(reading_date, row.get('notes')):
                blocking_issues.append(
                    f'Показания строка {number}: дата {reading_date} указана при неизвестной точной дате, '
                    'но не соответствует согласованной технической схеме 5/20.'
                )
                continue
            technical_dated_readings += 1
        if not reading_date:
            if value is not None:
                undated_values += 1
                source = row.get('source_sheet') or 'не указан'
                source_row = row.get('source_row') or 'не указана'
                review_notes.append(
                    f'Показания строка {number}: {value} м³ без достоверной даты; '
                    f'сохранить как недатированное исходное значение в примечании счётчика '
                    f'(источник: {source}, строка {source_row}), Reading не создавать.'
                )
            continue

        if value is None:
            continue
        dated_readings += 1
        key = (meter_id, reading_date)
        if key in reading_keys:
            blocking_issues.append(f'Показания строка {number}: повтор meter_id + date.')
        reading_keys.add(key)

    issues = blocking_issues + review_notes
    return {
        'counts': counts,
        'issues': issues,
        'blocking_issues': blocking_issues,
        'review_notes': review_notes,
        'ready': not blocking_issues,
        'structure_ready': not blocking_issues,
        'reading_plan': {
            'dated': dated_readings,
            'technical_dated': technical_dated_readings,
            'undated_to_meter_notes': undated_values,
            'empty_skipped': empty_readings,
        },
        'summary': (
            f'Участков: {len(plot_ids)}; людей: {len(person_ids)}; узлов: {len(node_names)}; '
            f'групп: {len(group_names)}; счётчиков: {len(meter_ids)}; '
            f'строк показаний: {counts["Показания"]}; блокирующих ошибок: {len(blocking_issues)}; '
            f'требуют сохранения/проверки: {len(review_notes)}.'
        ),
    }
