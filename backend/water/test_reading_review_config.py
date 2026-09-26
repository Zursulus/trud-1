from datetime import date
from io import BytesIO

from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from .models import Account, Meter, Reading, SupplyNode, User, WaterGroup
from .reading_review_config import export_readings_xlsx, reading_review_view


class ConfigurableReadingReviewTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='threshold-admin', email='threshold@example.test', password='test-pass-123',
        )
        self.node = SupplyNode.objects.create(name='Узел порога')
        self.group = WaterGroup.objects.create(name='Тестовая линия', node=self.node)
        self.account = Account.objects.create(number='TRUD-PLOT-T', plot='Тестовый участок')
        self.meter = Meter.objects.create(
            serial='TRUD-PLOT-T', kind='individual', node=self.node, group=self.group,
            account=self.account, notes='ID импорта: IND-TRUD-PLOT-T',
        )
        Reading.objects.create(meter=self.meter, date=date(2026, 6, 20), value='100')
        self.delta19 = Reading.objects.create(
            meter=self.meter, date=date(2026, 7, 20), value='119',
        )
        self.delta20 = Reading.objects.create(
            meter=self.meter, date=date(2026, 8, 20), value='139',
        )

    def request(self, path, params=None):
        request = RequestFactory().get(path, data=params or {}, HTTP_HOST='testserver')
        request.user = self.admin
        request.user.is_verified = lambda: True
        return request

    def test_default_threshold_is_20_and_inclusive(self):
        response = reading_review_view(self.request('/admin/water/readings/review/'))
        response.render()
        self.assertEqual(response.context_data['threshold'], '20')
        self.assertEqual(response.context_data['review_count'], 1)
        self.assertEqual(response.context_data['rows'][0]['reading'].pk, self.delta20.pk)
        html = response.content.decode('utf-8')
        self.assertIn('Порог индивидуального расхода, м³', html)
        self.assertIn('value="20"', html)
        self.assertIn('threshold=20', html)
        self.assertIn('Расход 20 м³ или больше', html)

    def test_admin_can_raise_or_lower_threshold_without_code_change(self):
        response = reading_review_view(
            self.request('/admin/water/readings/review/', {'threshold': '25'})
        )
        response.render()
        self.assertEqual(response.context_data['threshold'], '25')
        self.assertEqual(response.context_data['review_count'], 0)

        response = reading_review_view(
            self.request('/admin/water/readings/review/', {'threshold': '15'})
        )
        response.render()
        self.assertEqual(response.context_data['threshold'], '15')
        self.assertEqual(response.context_data['review_count'], 2)

    def test_invalid_threshold_falls_back_to_20_with_visible_error(self):
        response = reading_review_view(
            self.request('/admin/water/readings/review/', {'threshold': 'не число'})
        )
        response.render()
        self.assertEqual(response.context_data['threshold'], '20')
        self.assertEqual(response.context_data['review_count'], 1)
        self.assertIn('Введите число', response.content.decode('utf-8'))

    def test_xlsx_uses_same_selected_threshold(self):
        response = export_readings_xlsx(
            self.request('/admin/water/readings/export-xlsx/', {'threshold': '25'})
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook['Сводка']['B4'].value, 0)
        self.assertIn('≥ 25', workbook['Сводка']['B7'].value)
        self.assertEqual(workbook['Требуют проверки'].max_row, 1)

        response = export_readings_xlsx(
            self.request('/admin/water/readings/export-xlsx/', {'threshold': '20'})
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook['Сводка']['B4'].value, 1)
        self.assertIn('≥ 20', workbook['Сводка']['B7'].value)

    def test_invalid_xlsx_threshold_is_rejected(self):
        response = export_readings_xlsx(
            self.request('/admin/water/readings/export-xlsx/', {'threshold': 'oops'})
        )
        self.assertEqual(response.status_code, 400)
