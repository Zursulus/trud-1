import csv
import hashlib
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q

from .models import Account, ImportBatch, ImportRow, LandPlot, Person, PlotRelation

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024
MAX_ROWS = 5000

HEADER_ALIASES = {
    'account_number': {'лицевой счет', 'лицевой счёт', 'л с', 'лс', 'account', 'account number'},
    'plot_label': {'участок', 'номер участка', 'уч', 'plot', 'plot number'},
    'address': {'адрес', 'ориентир', 'address'},
    'cadastral_number': {'кадастровый номер', 'кадастровый', 'кадастр', 'cadastral number'},
    'area_m2': {'площадь', 'площадь м2', 'площадь м²', 'area', 'area m2'},
    'person_name': {'фио', 'ф и о', 'владелец', 'собственник', 'фамилия имя отчество', 'full name'},
    'phone': {'телефон', 'тел', 'phone'},
    'email': {'email', 'e mail', 'электронная почта', 'почта'},
}


def normalized_header(value):
    text = str(value or '').strip().lower().replace('ё', 'е')
    return re.sub(r'[^a-zа-я0-9]+', ' ', text).strip()


def header_map(values):
    aliases = {normalized_header(alias): field for field, items in HEADER_ALIASES.items() for alias in items}
    result = {}
    for index, value in enumerate(values):
        field = aliases.get(normalized_header(value))
        if field and field not in result:
            result[field] = index
    return result


def text_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_csv(data):
    decoded = None
    for encoding in ('utf-8-sig', 'cp1251'):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValidationError('CSV должен быть в UTF-8 или Windows-1251.')
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=';,\t,')
    except csv.Error:
        return 'CSV', list(csv.reader(io.StringIO(decoded), delimiter=';'))
    return 'CSV', list(csv.reader(io.StringIO(decoded), dialect))


def read_xlsx(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_SIZE:
                raise ValidationError('Распакованный Excel слишком велик.')
    except zipfile.BadZipFile as error:
        raise ValidationError('Файл не является корректным XLSX.') from error
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as error:
        raise ValidationError('Не удалось безопасно прочитать XLSX.') from error
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 2)]
    workbook.close()
    return sheet.title, rows


def parsed_rows(filename, data):
    lower = filename.lower()
    if lower.endswith('.csv'):
        return read_csv(data)
    if lower.endswith('.xlsx'):
        return read_xlsx(data)
    raise ValidationError('Поддерживаются только CSV и XLSX без макросов.')


@transaction.atomic
def stage_import(upload, effective_date, *, actor=None):
    data = upload.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE:
        raise ValidationError('Файл больше 5 МБ.')
    digest = hashlib.sha256(data).hexdigest()
    if ImportBatch.objects.filter(sha256=digest).exists():
        raise ValidationError('Этот файл уже загружался; повторный пакет не создан.')
    sheet, rows = parsed_rows(upload.name, data)
    if not rows:
        raise ValidationError('Файл пуст.')
    mapping = header_map(rows[0])
    if not ({'account_number', 'plot_label'} & set(mapping)):
        raise ValidationError('Не найдена колонка лицевого счёта или участка.')
    batch = ImportBatch(
        filename=upload.name[:255], sha256=digest, sheet=sheet[:200], effective_date=effective_date,
    )
    batch._history_user = actor
    batch._change_reason = 'Предварительная загрузка файла'
    batch.save()
    seen = set()
    count = 0
    for row_number, values in enumerate(rows[1:MAX_ROWS + 1], 2):
        if not any(value not in (None, '') for value in values):
            continue
        payload = {field: text_value(values[index]) if index < len(values) else '' for field, index in mapping.items()}
        issues = []
        if any(value.startswith('=') for value in payload.values()):
            issues.append('Обнаружена формула; строка требует ручной проверки.')
        area = None
        if payload.get('area_m2'):
            try:
                area = Decimal(payload['area_m2'].replace(',', '.'))
                if area <= 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                issues.append('Некорректная площадь.')
                area = None
        if not payload.get('account_number') and not payload.get('plot_label'):
            issues.append('Нет лицевого счёта и участка.')
        key = (payload.get('account_number', '').casefold(), payload.get('plot_label', '').casefold())
        if key in seen:
            issues.append('Повтор лицевого счёта и участка внутри файла.')
        seen.add(key)
        import_row = ImportRow(
            batch=batch, row_number=row_number, account_number=payload.get('account_number', ''),
            plot_label=payload.get('plot_label', ''), address=payload.get('address', ''),
            cadastral_number=payload.get('cadastral_number', ''), area_m2=area,
            person_name=payload.get('person_name', ''), phone=payload.get('phone', ''),
            email=payload.get('email', ''), status='review' if issues else 'ready',
            issues=' '.join(issues),
        )
        import_row._history_user = actor
        import_row._change_reason = 'Распознавание строки импорта'
        try:
            import_row.save()
        except ValidationError as error:
            import_row.status = 'review'
            import_row.email = ''
            import_row.issues = f'{import_row.issues} {error}'.strip()
            import_row.save()
        count += 1
    batch.row_count = count
    batch._history_user = actor
    batch._change_reason = 'Завершение предварительной проверки'
    batch.save()
    return batch


@transaction.atomic
def apply_import_row(row, *, actor=None):
    row = ImportRow.objects.select_for_update().select_related('batch').get(pk=row.pk)
    if row.status != 'ready':
        raise ValidationError('Применять можно только проверенную строку со статусом «Готово».')
    account = Account.objects.filter(number=row.account_number).first() if row.account_number else None
    if account is None:
        account = Account(number=row.account_number or None, plot=row.plot_label, contact_name=row.person_name, phone=row.phone)
        account._history_user = actor
        account._change_reason = f'Импорт из {row.batch.filename}, строка {row.row_number}'
        account.save()
    plot = None
    candidates = LandPlot.objects.none()
    if row.cadastral_number:
        candidates = LandPlot.objects.filter(cadastral_number=row.cadastral_number)
    elif row.plot_label:
        candidates = LandPlot.objects.filter(label=row.plot_label)
    if candidates.count() > 1:
        raise ValidationError('Найдено несколько подходящих участков; нужна ручная проверка.')
    plot = candidates.first()
    if plot and plot.account_id not in (None, account.pk):
        raise ValidationError('Найденный участок уже связан с другим лицевым счётом.')
    if plot and plot.account_id is None:
        plot.account = account
        plot._history_user = actor
        plot._change_reason = f'Связь со счётом по импорту {row.batch.filename}, строка {row.row_number}'
        plot.save()
    if plot is None and row.plot_label:
        plot = LandPlot(
            label=row.plot_label, address=row.address, cadastral_number=row.cadastral_number or None,
            area_m2=row.area_m2, account=account,
        )
        plot._history_user = actor
        plot._change_reason = f'Импорт из {row.batch.filename}, строка {row.row_number}'
        plot.save()
    person = None
    if row.person_name:
        lookup = Q(full_name=row.person_name, phone=row.phone)
        if row.email:
            lookup |= Q(email=row.email)
        matches = Person.objects.filter(lookup)
        if matches.count() > 1:
            raise ValidationError('Найдено несколько подходящих людей; нужна ручная проверка.')
        person = matches.first()
        if person is None:
            person = Person(full_name=row.person_name, phone=row.phone, email=row.email)
            person._history_user = actor
            person._change_reason = f'Импорт из {row.batch.filename}, строка {row.row_number}'
            person.save()
    if person and plot and not PlotRelation.objects.filter(
        person=person, plot=plot, role='owner', starts__lte=row.batch.effective_date,
    ).filter(Q(ends__isnull=True) | Q(ends__gt=row.batch.effective_date)).exists():
        relation = PlotRelation(person=person, plot=plot, role='owner', starts=row.batch.effective_date, document='Импорт; требует сверки с документами')
        relation._history_user = actor
        relation._change_reason = f'Импорт из {row.batch.filename}, строка {row.row_number}'
        relation.save()
    row.status = 'applied'
    row.applied_account, row.applied_plot, row.applied_person = account, plot, person
    row._history_user = actor
    row._change_reason = 'Применение проверенной строки импорта'
    row.save()
    batch = row.batch
    remaining = batch.rows.exclude(status__in=('applied', 'skipped')).exists()
    batch.status = 'partial' if remaining else 'applied'
    batch._history_user = actor
    batch._change_reason = 'Обновление состояния импорта'
    batch.save()
    return row
