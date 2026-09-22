from datetime import date
from io import StringIO
import tempfile

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from .models import Account, Membership, SupplyNode, WaterGroup


class CurrentWaterMembershipImportTests(TestCase):
    def setUp(self):
        self.node = SupplyNode.objects.create(name='Узел')
        self.group = WaterGroup.objects.create(name='Линия', node=self.node)
        self.other_group = WaterGroup.objects.create(name='Другая линия', node=self.node)
        self.account = Account.objects.create(number='TRUD-PLOT-0001', plot='Участок 1')

    def package(self, *, starts=''):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Участки'
        ws.append(['plot_id','label','address','cadastral_number','contact_name','phone','source_sheet','source_row','status'])
        ws.append(['TRUD-PLOT-0001','Участок 1','Участок 1','','','','Исходник',1,'готово'])

        ws = wb.create_sheet('Люди')
        ws.append(['person_id','full_name','phone','plot_id','role','notes'])

        ws = wb.create_sheet('Узлы')
        ws.append(['node_name','notes'])
        ws.append(['Узел',''])

        ws = wb.create_sheet('Группы')
        ws.append(['group_name','node_name','source','notes'])
        ws.append(['Линия','Узел','unknown',''])

        ws = wb.create_sheet('Состав групп')
        ws.append(['plot_id','group_name','starts','ends','notes'])
        ws.append([
            'TRUD-PLOT-0001','Линия',starts,'',
            'Группа определена по исходному листу; точная дата начала неизвестна',
        ])

        ws = wb.create_sheet('Индивидуальные счетчики')
        ws.append(['meter_id','kind','node_name','group_name','plot_id','serial','notes'])

        ws = wb.create_sheet('Общие и контрольные')
        ws.append(['meter_id','kind','node_name','group_name','serial','source_text','source_sheet','source_row','status'])

        ws = wb.create_sheet('Показания')
        ws.append(['meter_id','date','value_m3','notes','source_sheet','source_row'])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            path = tmp.name
        wb.save(path)
        wb.close()
        return path

    def test_dry_run_does_not_write(self):
        out = StringIO()
        call_command(
            'import_water_current_memberships', self.package(),
            effective_date='2026-09-22', stdout=out,
        )
        self.assertEqual(Membership.objects.count(), 0)
        self.assertIn('Новых Membership: 1', out.getvalue())
        self.assertIn('Историческая дата вступления НЕ утверждается', out.getvalue())

    def test_apply_uses_registry_effective_date_only(self):
        call_command(
            'import_water_current_memberships', self.package(),
            effective_date='2026-09-22', apply=True,
            confirm='IMPORT-CURRENT-WATER-MEMBERSHIPS', stdout=StringIO(),
        )
        membership = Membership.objects.get()
        self.assertEqual(membership.starts, date(2026, 9, 22))
        self.assertEqual(membership.group, self.group)
        history = membership.history.order_by('-history_date', '-history_id').first()
        self.assertIn('историческая дата начала неизвестна', history.history_change_reason)

    def test_source_historical_date_is_rejected(self):
        with self.assertRaises(CommandError):
            call_command(
                'import_water_current_memberships', self.package(starts=date(2026, 1, 1)),
                effective_date='2026-09-22', stdout=StringIO(),
            )
        self.assertEqual(Membership.objects.count(), 0)

    def test_conflicting_live_group_blocks_everything(self):
        Membership.objects.create(
            account=self.account, group=self.other_group,
            starts=date(2026, 9, 1), ends=None,
        )
        with self.assertRaises(CommandError):
            call_command(
                'import_water_current_memberships', self.package(),
                effective_date='2026-09-22', apply=True,
                confirm='IMPORT-CURRENT-WATER-MEMBERSHIPS', stdout=StringIO(),
            )
        self.assertEqual(Membership.objects.count(), 1)
        self.assertEqual(Membership.objects.get().group, self.other_group)
