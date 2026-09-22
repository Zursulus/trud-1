from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import PermissionDenied, ValidationError
from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from .models import (
    Account, ControllerReadingSubmission, Meter, Reading, SupplyNode, User, WaterGroup,
)
from .reading_admin_tools import export_readings_xlsx, reassign_reading, reassign_reading_view


class ReadingAdminToolsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='rootadmin', email='root@example.test', password='test-pass-123',
        )
        self.staff = User.objects.create_user(
            username='operator', email='operator@example.test', password='test-pass-123',
            is_staff=True,
        )
        self.node = SupplyNode.objects.create(name='Узел 1')
        self.group = WaterGroup.objects.create(name='Морская', node=self.node)
        self.account55 = Account.objects.create(number='TRUD-PLOT-0055', plot='Морская 17')
        self.account56 = Account.objects.create(number='TRUD-PLOT-0056', plot='Морская 18')
        self.source = Meter.objects.create(
            serial='TRUD-PLOT-0055', kind='individual', node=self.node, group=self.group,
            account=self.account55, notes='ID импорта: IND-TRUD-PLOT-0055',
        )
        self.destination = Meter.objects.create(
            serial='TRUD-PLOT-0056', kind='individual', node=self.node, group=self.group,
            account=self.account56, notes='ID импорта: IND-TRUD-PLOT-0056',
        )
        Reading.objects.create(meter=self.source, date=date(2026, 8, 20), value='207')
        Reading.objects.create(meter=self.destination, date=date(2026, 8, 20), value='1177')
        self.reading = Reading.objects.create(
            meter=self.source, date=date(2026, 9, 21), value='1183',
            notes='Показание контролёра. Последние показания',
        )
        self.submission = ControllerReadingSubmission.objects.create(
            meter=self.source, date=date(2026, 9, 21), value='1183',
            notes='Последние показания', submitted_by=self.admin,
            status='approved', reading=self.reading,
        )

    def test_reassign_moves_reading_and_controller_submission_atomically(self):
        old_id = self.reading.pk
        replacement, returned_old_id, linked, old_label, new_label, full_reason = reassign_reading(
            reading_id=old_id,
            destination_meter_id=self.destination.pk,
            reason='Контролёр выбрал соседний участок',
            actor=self.admin,
            expected_version=self.reading.version,
        )

        self.assertEqual(returned_old_id, old_id)
        self.assertEqual(linked, 1)
        self.assertIn('0055', old_label)
        self.assertIn('0056', new_label)
        self.assertIn('Контролёр выбрал соседний участок', full_reason)
        self.assertFalse(Reading.objects.filter(pk=old_id).exists())
        self.assertEqual(replacement.meter, self.destination)
        self.assertEqual(replacement.date, date(2026, 9, 21))
        self.assertEqual(replacement.value, Decimal('1183'))

        self.submission.refresh_from_db()
        self.assertEqual(self.submission.meter, self.destination)
        self.assertEqual(self.submission.reading, replacement)

        history = replacement.history.filter(
            history_user=self.admin,
            history_change_reason__icontains='Исправление привязки',
        ).first()
        self.assertIsNotNone(history)
        self.assertLessEqual(len(history.history_change_reason), 100)
        self.assertTrue(
            Reading.history.filter(
                id=old_id, history_type='-', history_user=self.admin,
            ).exists()
        )

    def test_reassign_blocks_existing_reading_on_destination_date(self):
        Reading.objects.create(
            meter=self.destination, date=date(2026, 9, 21), value='1180',
        )
        with self.assertRaises(ValidationError):
            reassign_reading(
                reading_id=self.reading.pk,
                destination_meter_id=self.destination.pk,
                reason='Ошибка участка', actor=self.admin,
                expected_version=self.reading.version,
            )
        self.assertTrue(Reading.objects.filter(pk=self.reading.pk, meter=self.source).exists())
        self.submission.refresh_from_db()
        self.assertEqual(self.submission.meter, self.source)

    def test_reassign_blocks_sequence_break(self):
        previous = Reading.objects.get(meter=self.destination, date=date(2026, 8, 20))
        previous.value = '1200'
        previous._change_reason = 'Тест'
        previous.save()
        with self.assertRaises(ValidationError):
            reassign_reading(
                reading_id=self.reading.pk,
                destination_meter_id=self.destination.pk,
                reason='Ошибка участка', actor=self.admin,
                expected_version=self.reading.version,
            )
        self.assertTrue(Reading.objects.filter(pk=self.reading.pk, meter=self.source).exists())

    def test_reassign_requires_superuser(self):
        with self.assertRaises(PermissionDenied):
            reassign_reading(
                reading_id=self.reading.pk,
                destination_meter_id=self.destination.pk,
                reason='Ошибка участка', actor=self.staff,
                expected_version=self.reading.version,
            )

        request = RequestFactory().get('/admin/water/readings/1/reassign/')
        request.user = self.staff
        with self.assertRaises(PermissionDenied):
            reassign_reading_view(request, self.reading.pk)

    def test_xlsx_export_contains_audit_columns_and_review_flag(self):
        request = RequestFactory().get('/admin/water/readings/export-xlsx/')
        request.user = self.admin
        response = export_readings_xlsx(request)

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ['Показания', 'Пояснение'])
        sheet = workbook['Показания']
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers[0], 'Reading ID')
        self.assertIn('Import / meter ID', headers)
        self.assertIn('Контроль', headers)
        self.assertIn('Источник', headers)

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        row_by_id = {row[0]: row for row in rows}
        exported = row_by_id[self.reading.pk]
        self.assertEqual(exported[1], 'IND-TRUD-PLOT-0055')
        self.assertEqual(exported[5], 'Морская 17')
        self.assertEqual(exported[10], 976)
        self.assertIn('ПРОВЕРИТЬ', exported[11])
        self.assertIn('Контролёр:', exported[12])
