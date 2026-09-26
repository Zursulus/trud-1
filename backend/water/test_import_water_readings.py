from datetime import date
from io import StringIO
import tempfile

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook

from .models import Account, Membership, Meter, Reading, SupplyNode, WaterGroup


class FullWaterImportCommandTests(TestCase):
    def setUp(self):
        self.node = SupplyNode.objects.create(name='Узел')
        self.group = WaterGroup.objects.create(name='Линия', node=self.node, source='unknown')
        self.account = Account.objects.create(number='TRUD-PLOT-0001', plot='Участок 1')
        self.meter = Meter.objects.create(
            serial='IND-TRUD-PLOT-0001', kind='individual', node=self.node,
            account=self.account, group=self.group,
            notes='ID импорта: IND-TRUD-PLOT-0001',
        )

    def package(self, *, conflict=False):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Участки'
        ws.append(['plot_id','label','address','cadastral_number','contact_name','phone','source_sheet','source_row','status'])
        ws.append(['TRUD-PLOT-0001','Участок 1','Участок 1','','','','Исходник',1,'готово'])

        ws = wb.create_sheet('Люди')
        ws.append(['person_id','full_name','phone','plot_id','role','notes'])

        ws = wb.create_sheet('Узлы')
        ws.append(['node_name','notes'])
        ws.append(['Узел',''])

        ws = wb.create_sheet('Группы')
        ws.append(['group_name','node_name','source','notes'])
        ws.append(['Линия','Узел','individual',''])

        ws = wb.create_sheet('Состав групп')
        ws.append(['plot_id','group_name','starts','ends','notes'])
        ws.append(['TRUD-PLOT-0001','Линия',date(2026,1,1),'',''])

        ws = wb.create_sheet('Индивидуальные счетчики')
        ws.append(['meter_id','kind','node_name','group_name','plot_id','serial','notes','commissioned_on','retired_on'])
        ws.append([
            'IND-TRUD-PLOT-0001-H1','individual','Узел','Линия','TRUD-PLOT-0001',
            'HIST-1','исторический прибор','',date(2026,7,1),
        ])
        ws.append([
            'IND-TRUD-PLOT-0001','individual','Узел','Линия','TRUD-PLOT-0001',
            'IND-TRUD-PLOT-0001','текущий прибор',date(2026,7,1),'',
        ])

        ws = wb.create_sheet('Общие и контрольные')
        ws.append(['meter_id','kind','node_name','group_name','serial','source_text','source_sheet','source_row','status'])

        ws = wb.create_sheet('Показания')
        ws.append(['meter_id','date','value_m3','notes','source_sheet','source_row'])
        ws.append(['IND-TRUD-PLOT-0001-H1',date(2026,6,5),100,'Импорт из исходного Excel','Исходник',1])
        ws.append(['IND-TRUD-PLOT-0001-H1',date(2026,6,20),110,'Импорт из исходного Excel','Исходник',1])
        ws.append(['IND-TRUD-PLOT-0001',date(2026,7,5),0,'Импорт из исходного Excel','Исходник',1])
        ws.append(['IND-TRUD-PLOT-0001',date(2026,7,20),5,'Импорт из исходного Excel','Исходник',1])
        if conflict:
            ws.append(['IND-TRUD-PLOT-0001',date(2026,9,20),11,'Импорт из исходного Excel','Исходник',1])

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            path = tmp.name
        wb.save(path)
        wb.close()
        return path

    def test_dry_run_does_not_write(self):
        path = self.package()
        out = StringIO()
        call_command('import_water_readings', path, stdout=out)
        self.assertEqual(Membership.objects.count(), 0)
        self.assertEqual(Meter.objects.count(), 1)
        self.assertEqual(Reading.objects.count(), 0)
        self.group.refresh_from_db()
        self.assertEqual(self.group.source, 'unknown')
        self.assertIn('DRY-RUN', out.getvalue())

    def test_apply_creates_history_membership_and_group_source(self):
        path = self.package()
        out = StringIO()
        call_command(
            'import_water_readings', path, apply=True,
            confirm='IMPORT-DATED-WATER-READINGS', stdout=out,
        )
        self.assertEqual(Membership.objects.count(), 1)
        self.assertEqual(Meter.objects.count(), 2)
        self.assertEqual(Reading.objects.count(), 4)
        self.group.refresh_from_db()
        self.assertEqual(self.group.source, 'individual')
        self.meter.refresh_from_db()
        self.assertEqual(self.meter.commissioned_on, date(2026,7,1))
        hist = Meter.objects.get(notes__contains='ID импорта: IND-TRUD-PLOT-0001-H1')
        self.assertEqual(hist.retired_on, date(2026,7,1))
        self.assertIn('Импорт завершён', out.getvalue())

    def test_manual_same_value_is_preserved_and_skipped(self):
        Reading.objects.create(meter=self.meter, date=date(2026,7,20), value=5, notes='Внесено вручную')
        path = self.package()
        call_command(
            'import_water_readings', path, apply=True,
            confirm='IMPORT-DATED-WATER-READINGS', stdout=StringIO(),
        )
        reading = Reading.objects.get(meter=self.meter, date=date(2026,7,20))
        self.assertEqual(reading.notes, 'Внесено вручную')
        self.assertEqual(Reading.objects.filter(meter=self.meter).count(), 2)

    def test_manual_conflict_stops_everything(self):
        Reading.objects.create(meter=self.meter, date=date(2026,9,20), value=10, notes='Внесено вручную')
        path = self.package(conflict=True)
        with self.assertRaises(CommandError):
            call_command(
                'import_water_readings', path, apply=True,
                confirm='IMPORT-DATED-WATER-READINGS', stdout=StringIO(),
            )
        self.assertEqual(Membership.objects.count(), 0)
        self.assertEqual(Meter.objects.count(), 1)
        self.group.refresh_from_db()
        self.assertEqual(self.group.source, 'unknown')
        manual = Reading.objects.get(meter=self.meter, date=date(2026,9,20))
        self.assertEqual(manual.value, 10)
        self.assertEqual(manual.notes, 'Внесено вручную')
