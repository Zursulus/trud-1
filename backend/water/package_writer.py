from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .package_imports import inspect_water_package
from .models import Account, LandPlot, Meter, Person, SupplyNode, WaterGroup


def _rows(book, name):
    ws = book[name]
    header = [str(v or '').strip() for v in next(ws.iter_rows(values_only=True))]
    for values in ws.iter_rows(values_only=True):
        if not any(v not in (None, '') for v in values):
            continue
        yield {header[i]: values[i] if i < len(values) else None for i in range(len(header))}


def _text(value):
    return str(value or '').strip()


def _append_note(current, line):
    line = _text(line)
    if not line:
        return current or ''
    return f'{current}\n{line}'.strip() if current else line


@transaction.atomic
def import_verified_water_package(uploaded_file):
    """Import only facts represented by the verified workbook without inventing dates."""
    uploaded_file.seek(0)
    report = inspect_water_package(uploaded_file)
    if not report.get('structure_ready') or report.get('blocking_issues'):
        raise ValidationError('Пакет не прошёл предварительную проверку.')
    uploaded_file.seek(0)
    book = load_workbook(uploaded_file, read_only=True, data_only=True)

    guarded = (Account, LandPlot, Person, SupplyNode, WaterGroup, Meter)
    occupied = {m.__name__: m.objects.count() for m in guarded if m.objects.exists()}
    if occupied:
        raise ValidationError(f'Импорт разрешён только в пустую базу water: {occupied}')

    accounts = {}
    plots = {}
    for row in _rows(book, 'Участки'):
        plot_key = _text(row.get('plot_id'))
        label = _text(row.get('label')) or _text(row.get('address')) or plot_key
        address = _text(row.get('address'))
        source = _text(row.get('source_sheet'))
        source_row = _text(row.get('source_row'))
        status = _text(row.get('status'))
        source_note = f'Источник: {source}'
        if source_row:
            source_note += f'; строка: {source_row}'
        if status:
            source_note += f'; статус: {status}'
        account = Account.objects.create(
            number=plot_key,
            plot=label,
            contact_name=_text(row.get('contact_name')),
            phone=_text(row.get('phone')),
            notes=source_note,
        )
        plot = LandPlot.objects.create(label=label, address=address, account=account)
        accounts[plot_key] = account
        plots[plot_key] = plot

    people = {}
    for row in _rows(book, 'Люди'):
        key = _text(row.get('person_id'))
        people[key] = Person.objects.create(
            full_name=_text(row.get('full_name')),
            phone=_text(row.get('phone')),
            notes=_text(row.get('notes')),
        )

    nodes = {}
    for row in _rows(book, 'Узлы'):
        name = _text(row.get('node_name'))
        nodes[name] = SupplyNode.objects.create(name=name, notes=_text(row.get('notes')))

    groups = {}
    for row in _rows(book, 'Группы'):
        name = _text(row.get('group_name'))
        node_name = _text(row.get('node_name'))
        groups[name] = WaterGroup.objects.create(
            name=name,
            node=nodes[node_name],
            notes=_append_note(_text(row.get('notes')), f"Тип из источника: {_text(row.get('source'))}"),
        )

    meters = {}
    valid_kinds = {value for value, _label in Meter._meta.get_field('kind').choices}
    meter_rows = list(_rows(book, 'Индивидуальные счетчики')) + list(_rows(book, 'Общие и контрольные'))
    for row in meter_rows:
        key = _text(row.get('meter_id'))
        plot_key = _text(row.get('plot_id'))
        group_name = _text(row.get('group_name'))
        node_name = _text(row.get('node_name'))
        kind = _text(row.get('kind')) or 'individual'
        if kind not in valid_kinds:
            raise ValidationError(f'Неизвестное назначение счётчика {key}: {kind}')
        notes = _text(row.get('notes'))
        notes = _append_note(notes, f'ID импорта: {key}')
        source_text = _text(row.get('source_text'))
        source_sheet = _text(row.get('source_sheet'))
        source_row = _text(row.get('source_row'))
        status = _text(row.get('status'))
        if source_text:
            notes = _append_note(notes, f'Исходные данные: {source_text}')
        if source_sheet or source_row or status:
            notes = _append_note(notes, f'Источник: {source_sheet}; строка: {source_row}; статус: {status}')
        meter = Meter.objects.create(
            kind=kind,
            serial=_text(row.get('serial')) or key,
            account=accounts.get(plot_key),
            group=groups.get(group_name),
            node=nodes[node_name],
            notes=notes,
        )
        meters[key] = meter

    preserved = 0
    skipped = 0
    dated_deferred = 0
    for row in _rows(book, 'Показания'):
        meter = meters[_text(row.get('meter_id'))]
        raw_date = row.get('date')
        raw_value = row.get('value_m3')
        if raw_value in (None, ''):
            skipped += 1
            continue
        source = _text(row.get('source_sheet'))
        source_row = _text(row.get('source_row'))
        original_note = _text(row.get('notes'))
        if raw_date in (None, ''):
            line = f'Недатированное исходное показание: {_text(raw_value)}; источник: {source}; строка источника: {source_row}'
            if original_note:
                line += f'; примечание: {original_note}'
            meter.notes = _append_note(meter.notes, line)
            meter.save(update_fields=['notes'])
            preserved += 1
            continue
        line = f'Датированное исходное показание (не импортировано как Reading): {_text(raw_date)} = {_text(raw_value)}; источник: {source}'
        if source_row:
            line += f'; строка источника: {source_row}'
        if original_note:
            line += f'; примечание: {original_note}'
        meter.notes = _append_note(meter.notes, line)
        meter.save(update_fields=['notes'])
        dated_deferred += 1

    return {
        'accounts': Account.objects.count(),
        'plots': LandPlot.objects.count(),
        'people': Person.objects.count(),
        'nodes': SupplyNode.objects.count(),
        'groups': WaterGroup.objects.count(),
        'meters': Meter.objects.count(),
        'undated_values_preserved': preserved,
        'empty_values_skipped': skipped,
        'dated_values_deferred': dated_deferred,
        'plot_relations_created': 0,
        'memberships_created': 0,
    }
