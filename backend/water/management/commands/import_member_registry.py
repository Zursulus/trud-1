from collections import defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from water.models import Account
from water.private_registry import MemberRegistryEntry
from water.resident_numbers import ResidentNumberSlot


SHEET_NAME = 'Закрытый реестр'
BASE_EXPECTED_IDS = set(range(1, 301))
OPTIONAL_RESERVED_IDS = set(range(301, 311))
ALLOWED_IDS = BASE_EXPECTED_IDS | OPTIONAL_RESERVED_IDS
REQUIRED_COLUMNS = {
    '№ пользователя',
    'Телефон(ы) нормализованные',
    'Email',
    'Адрес участка',
    'Год вступления (точный)',
    'Источник года/основание',
    'Статус',
}


def _text(value):
    return str(value or '').strip()


def _int_or_none(value):
    if value in (None, ''):
        return None
    try:
        return int(value)
    except (TypeError, ValueError) as error:
        raise CommandError('В реестре найдено значение, которое должно быть целым числом.') from error


def _norm_address(value):
    text = _text(value).casefold()
    text = text.replace('ул.', '').replace('улица', '')
    for char in ('"', "'", '«', '»', '“', '”'):
        text = text.replace(char, '')
    return ' '.join(text.split())


def _read_registry(path):
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise CommandError('Не удалось открыть XLSX реестра.') from error
    if SHEET_NAME not in book.sheetnames:
        raise CommandError(f'В файле нет листа «{SHEET_NAME}». Используйте подготовленный XLSX.')
    sheet = book[SHEET_NAME]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as error:
        raise CommandError('Лист закрытого реестра пуст.') from error
    headers = [_text(value) for value in header_values]
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise CommandError('Не хватает колонок: ' + ', '.join(missing))
    positions = {name: headers.index(name) for name in REQUIRED_COLUMNS}

    result = []
    seen = set()
    for values in rows:
        if not any(value not in (None, '') for value in values):
            continue
        rid = _int_or_none(values[positions['№ пользователя']])
        if rid is None:
            raise CommandError('Есть строка без № пользователя.')
        if rid in seen:
            raise CommandError(f'Повторяется № пользователя {rid}.')
        seen.add(rid)
        result.append({
            'resident_number': rid,
            'phone': _text(values[positions['Телефон(ы) нормализованные']]),
            'email': _text(values[positions['Email']]).lower(),
            'address': _text(values[positions['Адрес участка']]),
            'joined_year': _int_or_none(values[positions['Год вступления (точный)']]),
            'membership_note': _text(values[positions['Источник года/основание']]),
            'status': _text(values[positions['Статус']]),
        })

    missing_ids = sorted(BASE_EXPECTED_IDS - seen)
    unsupported_ids = sorted(seen - ALLOWED_IDS)
    if missing_ids or unsupported_ids:
        details = []
        if missing_ids:
            details.append('нет обязательных №: ' + ', '.join(map(str, missing_ids[:20])))
        if unsupported_ids:
            details.append('недопустимые №: ' + ', '.join(map(str, unsupported_ids[:20])))
        raise CommandError(
            'Ожидались обязательные №1–300; дополнительно разрешены только резервные №301–310; '
            + '; '.join(details)
        )
    return sorted(result, key=lambda row: row['resident_number'])


def _account_index():
    index = defaultdict(list)
    for account in Account.objects.filter(archived=False).only('id', 'number', 'plot'):
        key = _norm_address(account.plot)
        if key:
            index[key].append(account)
    return index


class Command(BaseCommand):
    help = (
        'Проверяет или импортирует закрытый реестр №1–300 без ФИО; '
        'при необходимости принимает резервные №301–310. '
        'По умолчанию выполняет только dry-run; запись требует --apply.'
    )

    def add_arguments(self, parser):
        parser.add_argument('xlsx', help='Путь к подготовленному XLSX')
        parser.add_argument('--apply', action='store_true', help='Записать закрытый реестр после успешной проверки')

    def handle(self, *args, **options):
        path = Path(options['xlsx']).expanduser()
        if not path.is_file():
            raise CommandError('Файл реестра не найден.')
        rows = _read_registry(path)
        row_ids = {row['resident_number'] for row in rows}

        slots = {
            slot.number: slot
            for slot in ResidentNumberSlot.objects.filter(number__in=row_ids)
        }
        if set(slots) != row_ids:
            missing_slots = sorted(row_ids - set(slots))
            raise CommandError(
                'Не подготовлены resident slots: ' + ', '.join(map(str, missing_slots))
            )
        bad_slots = [
            number for number, slot in slots.items()
            if slot.purpose != ResidentNumberSlot.PURPOSE_RESIDENT
        ]
        if bad_slots:
            raise CommandError('Некорректное назначение resident slots: ' + ', '.join(map(str, bad_slots)))

        accounts = _account_index()
        single_matches = {}
        ambiguous_ids = []
        missing_account_ids = []
        manual_source_ids = []
        for row in rows:
            rid = row['resident_number']
            if row['status'] != 'ГОТОВО':
                manual_source_ids.append(rid)
            candidates = accounts.get(_norm_address(row['address']), []) if row['address'] else []
            if len(candidates) == 1:
                single_matches[rid] = candidates[0]
            elif len(candidates) > 1:
                ambiguous_ids.append(rid)
            else:
                missing_account_ids.append(rid)

        extra_ids = sorted(row_ids - BASE_EXPECTED_IDS)
        self.stdout.write(
            f'Реестр: {len(rows)} записей; обязательные №1–300 присутствуют; ФИО не требуются и не создаются.'
        )
        if extra_ids:
            self.stdout.write('Дополнительные резервные №: ' + ', '.join(map(str, extra_ids)))
        self.stdout.write(f'Однозначно связать с действующим Account: {len(single_matches)}')
        self.stdout.write(f'Неоднозначный Account по адресу: {len(ambiguous_ids)}')
        self.stdout.write(f'Пока без действующего Account: {len(missing_account_ids)}')
        self.stdout.write(f'Строки источника для ручной проверки: {len(manual_source_ids)}')
        if ambiguous_ids:
            self.stdout.write('Неоднозначные №: ' + ', '.join(map(str, ambiguous_ids)))
        if manual_source_ids:
            self.stdout.write('Ручная проверка №: ' + ', '.join(map(str, manual_source_ids)))
        self.stdout.write('Контакты и иные значения ПД намеренно не выводятся.')

        if not options['apply']:
            self.stdout.write(self.style.WARNING('DRY-RUN: база данных не изменена. Для записи повторите с --apply.'))
            return

        created = 0
        updated = 0
        linked = 0
        with transaction.atomic():
            for row in rows:
                rid = row['resident_number']
                desired_account = single_matches.get(rid)
                entry = MemberRegistryEntry.objects.select_for_update().filter(pk=rid).first()
                if entry is None:
                    entry = MemberRegistryEntry(
                        resident_number=slots[rid],
                        account=desired_account,
                        phone=row['phone'],
                        email=row['email'],
                        joined_year=row['joined_year'],
                        membership_note=row['membership_note'],
                    )
                    entry.save()
                    created += 1
                    linked += int(bool(desired_account))
                    continue

                changes = []
                if entry.person_id:
                    # Legal/legacy Person is preserved, but never required or created by this import.
                    pass
                if desired_account:
                    if entry.account_id is None:
                        entry.account = desired_account
                        changes.append('account')
                        linked += 1
                    elif entry.account_id != desired_account.id:
                        raise CommandError(
                            f'№{rid} уже связан с другим Account. Автоматическая перепривязка запрещена.'
                        )

                for field in ('phone', 'email', 'joined_year', 'membership_note'):
                    desired = row[field]
                    current = getattr(entry, field)
                    empty = current in (None, '')
                    if empty and desired not in (None, ''):
                        setattr(entry, field, desired)
                        changes.append(field)
                    elif desired not in (None, '') and current != desired:
                        raise CommandError(
                            f'№{rid}: закрытые данные уже отличаются от файла. '
                            'Автоматическая перезапись запрещена; сначала проверьте запись вручную.'
                        )
                if changes:
                    entry.save()
                    updated += 1

        self.stdout.write(self.style.SUCCESS(
            f'Готово: создано {created}, дополнено {updated}, новых связей с Account {linked}. '
            'User-учётки и Person не создавались; №333 не затронут.'
        ))
