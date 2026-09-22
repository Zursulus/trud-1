from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from water.models import Meter, Reading
from water.package_imports import inspect_water_package


def _text(value):
    return str(value or '').strip()


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for parser in (
        lambda s: date.fromisoformat(s),
        lambda s: datetime.strptime(s, '%d.%m.%Y').date(),
    ):
        try:
            return parser(text)
        except (ValueError, TypeError):
            pass
    raise CommandError(f'Не удалось распознать дату показания: {text or "(пусто)"}')


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(v) for v in next(values))
    except StopIteration:
        return
    for values in values:
        if not any(v not in (None, '') for v in values):
            continue
        normalized = tuple(_text(v) for v in values)
        if normalized[:len(header)] == header:
            continue
        yield {header[i]: values[i] if i < len(values) else None for i in range(len(header))}


class Command(BaseCommand):
    help = 'Безопасно добавить датированные показания из проверенного XLSX-пакета в существующий water-реестр.'

    def add_arguments(self, parser):
        parser.add_argument('file', type=Path)
        parser.add_argument('--apply', action='store_true', help='Выполнить запись. Без флага только dry-run.')
        parser.add_argument('--confirm', default='')

    def handle(self, *args, **options):
        path = options['file'].resolve()
        if not path.is_file():
            raise CommandError(f'Файл не найден: {path}')

        raw = path.read_bytes()
        upload = SimpleUploadedFile(path.name, raw)
        report = inspect_water_package(upload)
        if report.get('blocking_issues'):
            raise CommandError('Пакет не прошёл проверку: ' + '; '.join(report['blocking_issues']))

        book = load_workbook(path, read_only=True, data_only=True)
        prepared = []
        missing_meters = []
        conflicts = []

        for row in _rows(book['Показания']):
            raw_date = row.get('date')
            raw_value = row.get('value_m3')
            if raw_date in (None, '') or raw_value in (None, ''):
                continue

            meter_id = _text(row.get('meter_id'))
            matches = list(Meter.objects.filter(notes__contains=f'ID импорта: {meter_id}').order_by('pk')[:2])
            if len(matches) != 1:
                missing_meters.append(f'{meter_id}: найдено {len(matches)} счётчиков')
                continue
            meter = matches[0]
            reading_date = _parse_date(raw_date)
            try:
                value = Decimal(str(raw_value).replace(',', '.'))
            except (InvalidOperation, ValueError) as error:
                raise CommandError(f'{meter_id} {reading_date}: некорректное значение {raw_value}') from error

            existing = Reading.objects.filter(meter=meter, date=reading_date).first()
            if existing:
                if existing.value != value:
                    conflicts.append(
                        f'{meter_id} {reading_date}: в базе {existing.value}, в пакете {value}'
                    )
                prepared.append((meter_id, meter, reading_date, value, row, True))
                continue
            prepared.append((meter_id, meter, reading_date, value, row, False))

        book.close()

        if missing_meters:
            raise CommandError('Неоднозначная привязка счётчиков: ' + '; '.join(missing_meters[:20]))
        if conflicts:
            raise CommandError('Есть конфликтующие существующие показания: ' + '; '.join(conflicts[:20]))

        to_create = [item for item in prepared if not item[5]]
        already_same = len(prepared) - len(to_create)
        self.stdout.write(f'Проверено датированных строк: {len(prepared)}')
        self.stdout.write(f'Будет создано: {len(to_create)}')
        self.stdout.write(f'Уже есть с тем же значением: {already_same}')

        if not options['apply']:
            self.stdout.write(self.style.WARNING('DRY-RUN: база не изменена. Для записи добавьте --apply и подтверждение.'))
            return
        if options['confirm'] != 'IMPORT-DATED-WATER-READINGS':
            raise CommandError('Для записи нужен --confirm IMPORT-DATED-WATER-READINGS')

        created = 0
        try:
            with transaction.atomic():
                for meter_id, meter, reading_date, value, row, existed in sorted(
                    to_create, key=lambda item: (item[1].pk, item[2])
                ):
                    source = _text(row.get('source_sheet'))
                    source_row = _text(row.get('source_row'))
                    original_note = _text(row.get('notes'))
                    note = f'Импорт из подготовленного пакета; источник: {source}; строка: {source_row}'
                    if original_note:
                        note += f'; {original_note}'
                    Reading.objects.create(
                        meter=meter,
                        date=reading_date,
                        value=value,
                        notes=note,
                    )
                    created += 1
        except ValidationError as error:
            raise CommandError(f'Импорт отменён целиком: {error}') from error

        self.stdout.write(self.style.SUCCESS(f'Импорт завершён: создано {created} показаний.'))
