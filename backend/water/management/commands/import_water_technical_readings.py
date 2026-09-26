from calendar import monthrange
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from water.models import Account, Meter, Reading, SupplyNode, WaterGroup


CONFIRM = 'IMPORT-TECHNICAL-WATER-READINGS'
MARKER = 'техническая дата по согласованной схеме граница/20'
MONTHS = {'июнь': 6, 'июль': 7, 'август': 8}
MAX_FILE_SIZE = 5 * 1024 * 1024


def _text(value):
    return str(value or '').strip()


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(v) for v in next(values))
    except StopIteration:
        return
    for values in values:
        if not any(v not in (None, '') for v in values):
            continue
        normalized = tuple(_text(v) for v in values)
        if normalized[:len(header)] == header:
            continue
        yield {header[i]: values[i] if i < len(values) else None for i in range(len(header))}


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for parser in (
        lambda s: date.fromisoformat(s[:10]),
        lambda s: datetime.strptime(s, '%d.%m.%Y').date(),
    ):
        try:
            return parser(text)
        except (ValueError, TypeError):
            pass
    raise CommandError(f'Не удалось распознать дату: {text or "(пусто)"}')


def _import_meter_id(meter):
    marker = 'ID импорта:'
    for line in _text(meter.notes).splitlines():
        if marker in line:
            return line.split(marker, 1)[1].strip().split(';', 1)[0].strip()
    return ''


def _technical_date_is_valid(reading_date, note, meter_id):
    low = _text(note).casefold()
    if MARKER not in low or 'точн' not in low or 'дат' not in low or 'неизвест' not in low:
        return False
    month = next((number for name, number in MONTHS.items() if name in low), None)
    if month is None:
        return False

    if 'предыдущее' in low or 'начальное' in low:
        if month == 1:
            expected = date(reading_date.year - 1, 12, 31)
        else:
            previous_month = month - 1
            expected = date(reading_date.year, previous_month, monthrange(reading_date.year, previous_month)[1])
        return reading_date == expected

    if 'текущее' in low:
        return reading_date == date(reading_date.year, month, 20)

    # В исходнике ровно одна строка без роли: августовское значение SYS-METER-012.
    return meter_id == 'SYS-METER-012' and reading_date == date(reading_date.year, 8, 20)


class Command(BaseCommand):
    help = (
        'Импортирует только исторические показания с явно помеченными техническими датами: '
        'предыдущее/начальное — последний день предыдущего месяца, текущее — 20-е число. '
        'Существующие ручные показания никогда не перезаписываются.'
    )

    def add_arguments(self, parser):
        parser.add_argument('file', type=Path)
        parser.add_argument('--apply', action='store_true')
        parser.add_argument('--confirm', default='')

    def handle(self, *args, **options):
        path = options['file'].resolve()
        if not path.is_file():
            raise CommandError(f'Файл не найден: {path}')
        if path.stat().st_size > MAX_FILE_SIZE:
            raise CommandError('Файл больше 5 МБ.')

        book = load_workbook(path, read_only=True, data_only=True)
        required = {'Показания', 'Индивидуальные счетчики'}
        missing = required - set(book.sheetnames)
        if missing:
            book.close()
            raise CommandError('Нет обязательных листов: ' + ', '.join(sorted(missing)))

        accounts = {a.number: a for a in Account.objects.exclude(number__isnull=True)}
        nodes = {n.name: n for n in SupplyNode.objects.all()}
        groups = {g.name: g for g in WaterGroup.objects.all()}

        db_meter_map = {}
        duplicate_ids = set()
        for meter in Meter.objects.select_related('account', 'node', 'group'):
            meter_id = _import_meter_id(meter)
            if not meter_id:
                continue
            if meter_id in db_meter_map:
                duplicate_ids.add(meter_id)
            db_meter_map[meter_id] = meter

        conflicts = []
        if duplicate_ids:
            conflicts.append('Повторяющиеся ID импорта счётчиков в базе: ' + ', '.join(sorted(duplicate_ids)))

        package_meters = {}
        for row in _rows(book['Индивидуальные счетчики']):
            meter_id = _text(row.get('meter_id'))
            if meter_id:
                package_meters[meter_id] = row

        meter_creates = []
        planned_ids = set(db_meter_map)
        for meter_id, row in package_meters.items():
            if meter_id in planned_ids or not meter_id.endswith('-H1'):
                continue
            plot_id = _text(row.get('plot_id'))
            node_name = _text(row.get('node_name'))
            group_name = _text(row.get('group_name'))
            account = accounts.get(plot_id)
            node = nodes.get(node_name)
            group = groups.get(group_name)
            if account is None or node is None or group is None:
                conflicts.append(f'{meter_id}: не удалось найти участок/узел/группу исторического прибора.')
                continue
            meter_creates.append({
                'meter_id': meter_id,
                'serial': _text(row.get('serial')) or meter_id,
                'kind': _text(row.get('kind')) or 'individual',
                'account': account,
                'node': node,
                'group': group,
                'notes': (
                    _text(row.get('notes')) + f'\nID импорта: {meter_id}'
                ).strip(),
            })
            planned_ids.add(meter_id)

        package = {}
        seen_keys = set()
        row_count = 0
        for row in _rows(book['Показания']):
            meter_id = _text(row.get('meter_id'))
            raw_date = row.get('date')
            raw_value = row.get('value_m3')
            note = _text(row.get('notes'))
            if not meter_id or raw_date in (None, '') or raw_value in (None, ''):
                conflicts.append('Показания: найдена неполная строка; v4 должен содержать meter_id/date/value.')
                continue
            row_count += 1
            if meter_id not in planned_ids:
                conflicts.append(f'{meter_id}: нет однозначного прибора для показания.')
                continue
            reading_date = _parse_date(raw_date)
            if not _technical_date_is_valid(reading_date, note, meter_id):
                conflicts.append(
                    f'{meter_id} {reading_date}: техническая дата не соответствует схеме граница/20.'
                )
                continue
            try:
                value = Decimal(str(raw_value).replace(',', '.'))
                if value < 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                conflicts.append(f'{meter_id} {reading_date}: некорректное значение {raw_value}.')
                continue
            key = (meter_id, reading_date)
            if key in seen_keys:
                conflicts.append(f'{meter_id} {reading_date}: повтор meter_id + date в пакете.')
                continue
            seen_keys.add(key)
            package.setdefault(meter_id, []).append((reading_date, value, note))

        book.close()

        to_create = []
        already_same = 0
        for meter_id, values in package.items():
            meter = db_meter_map.get(meter_id)
            existing_values = [] if meter is None else list(Reading.objects.filter(meter=meter).order_by('date', 'id'))
            existing_by_date = {reading.date: reading for reading in existing_values}
            combined = [(r.date, r.value, 'база') for r in existing_values]

            for reading_date, value, note in values:
                existing = existing_by_date.get(reading_date)
                if existing:
                    if existing.value == value:
                        already_same += 1
                    else:
                        conflicts.append(
                            f'{meter_id} {reading_date}: существующее значение {existing.value}, '
                            f'в пакете {value}; существующая запись имеет приоритет.'
                        )
                    continue
                combined.append((reading_date, value, 'пакет'))
                to_create.append((meter_id, reading_date, value, note))

            combined.sort(key=lambda item: item[0])
            previous = None
            for reading_date, value, origin in combined:
                if previous and value < previous[1]:
                    conflicts.append(
                        f'{meter_id}: последовательность ломается между {previous[0]}={previous[1]} '
                        f'и {reading_date}={value} ({origin}).'
                    )
                previous = (reading_date, value)

        if conflicts:
            unique = list(dict.fromkeys(conflicts))
            raise CommandError(
                f'DRY-RUN/импорт остановлен: найдено конфликтов {len(unique)}. '
                + ' | '.join(unique[:30])
            )

        self.stdout.write(f'Технических показаний в пакете: {row_count}')
        self.stdout.write(f'Новых показаний: {len(to_create)}; уже совпадают: {already_same}')
        self.stdout.write(f'Новых исторических счётчиков H1: {len(meter_creates)}')

        if not options['apply']:
            self.stdout.write(self.style.WARNING(
                f'DRY-RUN: база не изменена. Для записи добавьте --apply --confirm {CONFIRM}'
            ))
            return
        if options['confirm'] != CONFIRM:
            raise CommandError(f'Для записи нужен --confirm {CONFIRM}')

        with transaction.atomic():
            for spec in meter_creates:
                meter = Meter(
                    serial=spec['serial'], kind=spec['kind'], node=spec['node'],
                    account=spec['account'], group=spec['group'], notes=spec['notes'],
                )
                meter._change_reason = 'Импорт исторического прибора H1'
                meter.save()
                db_meter_map[spec['meter_id']] = meter

            created = 0
            for meter_id, reading_date, value, note in sorted(to_create, key=lambda x: (x[0], x[1])):
                meter = db_meter_map[meter_id]
                reading = Reading(
                    meter=meter,
                    date=reading_date,
                    value=value,
                    notes=note,
                )
                reading._change_reason = 'Импорт технически датированного исторического показания'
                reading.save()
                created += 1

        self.stdout.write(self.style.SUCCESS(
            f'Импорт завершён. Создано показаний: {created}; исторических H1: {len(meter_creates)}.'
        ))
