from datetime import date, datetime
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from water.models import Account, Membership, WaterGroup
from water.package_imports import inspect_water_package


CONFIRM = 'IMPORT-CURRENT-WATER-MEMBERSHIPS'


def _text(value):
    return str(value or '').strip()


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for parser in (
        lambda s: date.fromisoformat(s),
        lambda s: datetime.strptime(s, '%d.%m.%Y').date(),
    ):
        try:
            return parser(text)
        except (ValueError, TypeError):
            pass
    raise CommandError(f'Не удалось распознать дату: {text or "(пусто)"}')


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(v) for v in next(values))
    except StopIteration:
        return
    for values in values:
        if not any(v not in (None, '') for v in values):
            continue
        normalized = tuple(_text(v) for v in values)
        if normalized[:len(header)] == header:
            continue
        yield {header[i]: values[i] if i < len(values) else None for i in range(len(header))}


def _overlap(starts_a, ends_a, starts_b, ends_b):
    far = date.max
    return starts_a < (ends_b or far) and starts_b < (ends_a or far)


class Command(BaseCommand):
    help = (
        'Импортирует только текущий состав групп, когда историческая дата начала неизвестна. '
        'Переданная --effective-date означает дату, с которой связь разрешено использовать в реестре, '
        'и не считается исторической датой вступления в группу.'
    )

    def add_arguments(self, parser):
        parser.add_argument('file', type=Path)
        parser.add_argument(
            '--effective-date', required=True,
            help='Дата начала действия связи в реестре (YYYY-MM-DD); не историческая дата вступления.',
        )
        parser.add_argument('--apply', action='store_true', help='Выполнить запись. Без флага только dry-run.')
        parser.add_argument('--confirm', default='')

    def handle(self, *args, **options):
        path = options['file'].resolve()
        if not path.is_file():
            raise CommandError(f'Файл не найден: {path}')

        effective_date = _parse_date(options['effective_date'])
        if effective_date > timezone.localdate():
            raise CommandError('Дата действия состава групп не может быть в будущем.')

        raw = path.read_bytes()
        upload = SimpleUploadedFile(path.name, raw)
        report = inspect_water_package(upload)
        if report.get('blocking_issues'):
            raise CommandError('Пакет не прошёл проверку: ' + '; '.join(report['blocking_issues']))

        accounts = {account.number: account for account in Account.objects.exclude(number__isnull=True)}
        groups = {group.name: group for group in WaterGroup.objects.all()}

        book = load_workbook(path, read_only=True, data_only=True)
        planned = []
        already_current = 0
        source_rows = 0
        conflicts = []
        try:
            for row in _rows(book['Состав групп']):
                source_rows += 1
                plot_id = _text(row.get('plot_id'))
                group_name = _text(row.get('group_name'))
                raw_starts = row.get('starts')
                raw_ends = row.get('ends')
                note = _text(row.get('notes'))

                if raw_starts not in (None, '') or raw_ends not in (None, ''):
                    conflicts.append(
                        f'{plot_id}: команда текущего состава принимает только строки без исторических '
                        'starts/ends; исходные даты нужно разбирать отдельно.'
                    )
                    continue

                if 'дата' not in note.casefold() or 'неизвест' not in note.casefold():
                    conflicts.append(
                        f'{plot_id}: нет явной пометки, что историческая дата начала неизвестна.'
                    )
                    continue

                account = accounts.get(plot_id)
                group = groups.get(group_name)
                if account is None:
                    conflicts.append(f'{plot_id}: лицевой счёт/участок не найден в рабочей базе.')
                    continue
                if group is None:
                    conflicts.append(f'{plot_id}: группа «{group_name}» не найдена.')
                    continue

                overlaps = [
                    membership for membership in Membership.objects.filter(account=account).select_related('group')
                    if _overlap(effective_date, None, membership.starts, membership.ends)
                ]
                if overlaps:
                    same = [membership for membership in overlaps if membership.group_id == group.pk]
                    different = [membership for membership in overlaps if membership.group_id != group.pk]
                    if different:
                        current = ', '.join(
                            f'{membership.group.name} {membership.starts}–{membership.ends or "…"}'
                            for membership in overlaps
                        )
                        conflicts.append(
                            f'{plot_id}: с {effective_date} пакет подтверждает «{group_name}», '
                            f'но уже есть пересекающееся членство: {current}.'
                        )
                        continue
                    if same:
                        already_current += 1
                        continue

                planned.append((account, group))
        finally:
            book.close()

        if conflicts:
            unique = list(dict.fromkeys(conflicts))
            raise CommandError(
                f'DRY-RUN/импорт остановлен: найдено конфликтов {len(unique)}. '
                + ' | '.join(unique[:30])
            )

        self.stdout.write(f'Строк текущего состава в пакете: {source_rows}')
        self.stdout.write(f'Новых Membership: {len(planned)}; уже действуют в той же группе: {already_current}')
        self.stdout.write(
            f'Дата действия в реестре: {effective_date}. '
            'Историческая дата вступления НЕ утверждается и остаётся неизвестной.'
        )

        if not options['apply']:
            self.stdout.write(self.style.WARNING(
                f'DRY-RUN: база не изменена. Для записи добавьте --apply --confirm {CONFIRM}'
            ))
            return
        if options['confirm'] != CONFIRM:
            raise CommandError(f'Для записи нужен --confirm {CONFIRM}')

        created = 0
        with transaction.atomic():
            for account, group in planned:
                membership = Membership(
                    account=account,
                    group=group,
                    starts=effective_date,
                    ends=None,
                )
                membership._change_reason = (
                    'Текущий состав: группа подтверждена; историческая дата начала неизвестна.'
                )
                membership.save()
                created += 1

        self.stdout.write(self.style.SUCCESS(
            f'Импорт текущего состава завершён. Создано Membership: {created}. '
            'Исторические даты начала не создавались.'
        ))
