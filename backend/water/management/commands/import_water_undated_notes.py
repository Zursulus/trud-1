from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from water.models import Meter
from water.package_imports import inspect_water_package


CONFIRM = 'IMPORT-UNDATED-WATER-NOTES'


def _text(value):
    return str(value or '').strip()


def _rows(sheet):
    values = sheet.iter_rows(values_only=True)
    try:
        header = tuple(_text(v) for v in next(values))
    except StopIteration:
        return
    for values in values:
        if not any(v not in (None, '') for v in values):
            continue
        normalized = tuple(_text(v) for v in values)
        if normalized[:len(header)] == header:
            continue
        yield {header[i]: values[i] if i < len(values) else None for i in range(len(header))}


def _import_meter_id(meter):
    marker = 'ID импорта:'
    for line in _text(meter.notes).splitlines():
        if marker in line:
            return line.split(marker, 1)[1].strip().split(';', 1)[0].strip()
    return ''


def _note_line(row):
    raw_value = row.get('value_m3')
    source = _text(row.get('source_sheet'))
    source_row = _text(row.get('source_row'))
    original_note = _text(row.get('notes'))
    line = (
        f'Недатированное исходное показание: {_text(raw_value)}; '
        f'источник: {source}; строка источника: {source_row}'
    )
    if original_note:
        line += f'; примечание: {original_note}'
    return line


class Command(BaseCommand):
    help = (
        'Безопасно переносит недатированные исходные значения из проверенного XLSX '
        'в примечания существующих счётчиков. Reading не создаются.'
    )

    def add_arguments(self, parser):
        parser.add_argument('file', type=Path)
        parser.add_argument('--apply', action='store_true', help='Выполнить запись. Без флага только dry-run.')
        parser.add_argument('--confirm', default='')

    def handle(self, *args, **options):
        path = options['file'].resolve()
        if not path.is_file():
            raise CommandError(f'Файл не найден: {path}')

        raw = path.read_bytes()
        upload = SimpleUploadedFile(path.name, raw)
        report = inspect_water_package(upload)
        if report.get('blocking_issues'):
            raise CommandError('Пакет не прошёл проверку: ' + '; '.join(report['blocking_issues']))

        meters = {}
        duplicate_ids = set()
        for meter in Meter.objects.all():
            meter_id = _import_meter_id(meter)
            if not meter_id:
                continue
            if meter_id in meters:
                duplicate_ids.add(meter_id)
            meters[meter_id] = meter
        if duplicate_ids:
            raise CommandError('Повторяющиеся ID импорта счётчиков в базе: ' + ', '.join(sorted(duplicate_ids)))

        book = load_workbook(path, read_only=True, data_only=True)
        planned = {}
        already_same = 0
        conflicts = []
        source_values = 0
        try:
            for row in _rows(book['Показания']):
                raw_date = row.get('date')
                raw_value = row.get('value_m3')
                if raw_date not in (None, '') or raw_value in (None, ''):
                    continue
                source_values += 1
                meter_id = _text(row.get('meter_id'))
                meter = meters.get(meter_id)
                if meter is None:
                    conflicts.append(f'{meter_id}: счётчик не найден в рабочей базе.')
                    continue
                line = _note_line(row)
                existing_lines = set(_text(meter.notes).splitlines())
                pending_lines = planned.setdefault(meter.pk, [])
                if line in existing_lines or line in pending_lines:
                    already_same += 1
                    continue
                pending_lines.append(line)
        finally:
            book.close()

        if conflicts:
            unique = list(dict.fromkeys(conflicts))
            raise CommandError(
                f'DRY-RUN/импорт остановлен: найдено конфликтов {len(unique)}. '
                + ' | '.join(unique[:30])
            )

        additions = sum(len(lines) for lines in planned.values())
        self.stdout.write(f'Недатированных исходных значений в пакете: {source_values}')
        self.stdout.write(f'Новых строк в notes: {additions}; уже сохранены: {already_same}')

        if not options['apply']:
            self.stdout.write(self.style.WARNING(
                f'DRY-RUN: база не изменена. Для записи добавьте --apply --confirm {CONFIRM}'
            ))
            return
        if options['confirm'] != CONFIRM:
            raise CommandError(f'Для записи нужен --confirm {CONFIRM}')

        changed_meters = 0
        with transaction.atomic():
            for meter_pk, lines in planned.items():
                meter = Meter.objects.select_for_update().get(pk=meter_pk)
                current = _text(meter.notes)
                current_lines = set(current.splitlines())
                missing = [line for line in lines if line not in current_lines]
                if not missing:
                    continue
                meter.notes = '\n'.join([part for part in [current, *missing] if part]).strip()
                meter._change_reason = 'Импорт недатированных исходных показаний в примечания'
                meter.save(update_fields=['notes'])
                changed_meters += 1

        self.stdout.write(self.style.SUCCESS(
            f'Импорт завершён: добавлено строк {additions}, изменено счётчиков {changed_meters}. Reading не создавались.'
        ))
