from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import admin
from django.contrib.admin.models import LogEntry
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.template.response import TemplateResponse
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Reading
from .reading_admin_tools import (
    XLSX_HEADERS,
    _append_xlsx_row,
    _audit_readings,
    _style_data_sheet,
)


DEFAULT_REVIEW_THRESHOLD_M3 = Decimal('20')
MIN_REVIEW_THRESHOLD_M3 = Decimal('0.001')
MAX_REVIEW_THRESHOLD_M3 = Decimal('100000')


def _threshold_text(value):
    text = format(value, 'f')
    if '.' in text:
        text = text.rstrip('0').rstrip('.')
    return text


def _parse_threshold(raw_value):
    if raw_value is None or str(raw_value).strip() == '':
        return DEFAULT_REVIEW_THRESHOLD_M3, None
    try:
        value = Decimal(str(raw_value).strip().replace(',', '.'))
    except (InvalidOperation, ValueError):
        return DEFAULT_REVIEW_THRESHOLD_M3, 'Введите число, например 20 или 25.5.'
    if value < MIN_REVIEW_THRESHOLD_M3 or value > MAX_REVIEW_THRESHOLD_M3:
        return DEFAULT_REVIEW_THRESHOLD_M3, 'Порог должен быть больше 0 и не больше 100000 м³.'
    if value.as_tuple().exponent < -3:
        return DEFAULT_REVIEW_THRESHOLD_M3, 'Используйте не больше трёх знаков после запятой.'
    return value, None


def audit_readings_with_threshold(threshold):
    """Reuse canonical audit rules, replacing only the individual-consumption threshold."""
    result = []
    for original in _audit_readings():
        row = dict(original)
        reasons = [
            reason for reason in row['review_reason'].split('; ') if reason and not (
                reason.startswith('Расход ') and reason.endswith(' м³ или больше')
            )
        ]
        if (
            row['reading'].meter.kind == 'individual'
            and row['consumption'] is not None
            and row['consumption'] >= threshold
        ):
            reasons.append(f'Расход {_threshold_text(threshold)} м³ или больше')

        row['review_reason'] = '; '.join(reasons)
        row['review'] = bool(reasons)
        if original['severity'] == 'critical':
            row['severity'] = 'critical'
        elif reasons:
            row['severity'] = 'review'
        else:
            row['severity'] = 'ok'
        result.append(row)
    return result


def reading_review_view(request):
    if not request.user.has_perm('water.view_reading'):
        raise PermissionDenied

    threshold, threshold_error = _parse_threshold(request.GET.get('threshold'))
    rows = audit_readings_with_threshold(threshold)
    flagged = [row for row in rows if row['review']]
    threshold_text = _threshold_text(threshold)
    context = {
        **admin.site.each_context(request),
        'title': 'Проверка показаний счётчиков',
        'opts': Reading._meta,
        'rows': flagged,
        'total_count': len(rows),
        'review_count': len(flagged),
        'critical_count': sum(1 for row in flagged if row['severity'] == 'critical'),
        'controller_count': sum(1 for row in rows if row['controller_count']),
        'can_export': request.user.has_perm('water.export_reading'),
        'can_reassign': request.user.is_superuser,
        'threshold': threshold_text,
        'threshold_query': threshold_text,
        'threshold_error': threshold_error,
        'default_threshold': _threshold_text(DEFAULT_REVIEW_THRESHOLD_M3),
    }
    return TemplateResponse(request, 'admin/water/reading/review.html', context)


def export_readings_xlsx(request):
    if not request.user.has_perm('water.export_reading'):
        raise PermissionDenied

    threshold, threshold_error = _parse_threshold(request.GET.get('threshold'))
    if threshold_error:
        return HttpResponse(threshold_error, status=400, content_type='text/plain; charset=utf-8')

    rows = audit_readings_with_threshold(threshold)
    flagged = [row for row in rows if row['review']]
    controller_count = sum(1 for row in rows if row['controller_count'])
    imported_count = sum(
        1 for row in rows if row['source'] == 'Исторический импорт с технической датой'
    )
    threshold_text = _threshold_text(threshold)

    wb = Workbook()
    summary = wb.active
    summary.title = 'Сводка'
    summary.sheet_view.showGridLines = False
    summary.merge_cells('A1:D1')
    summary['A1'] = 'Проверка показаний ТСН «ТРУД-1»'
    summary['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    summary['A1'].fill = PatternFill('solid', fgColor='305496')
    summary['A1'].alignment = Alignment(vertical='center')
    summary.row_dimensions[1].height = 30
    summary.append(['Сформировано', timezone.localtime().strftime('%d.%m.%Y %H:%M'), '', ''])
    summary.append(['Всего показаний', len(rows), '', ''])
    summary.append(['Требуют проверки', len(flagged), '', ''])
    summary.append(['Связаны с контролёром', controller_count, '', ''])
    summary.append(['Исторический импорт', imported_count, '', ''])
    summary.append(['Порог контроля для индивидуального счётчика', f'≥ {threshold_text} м³', '', ''])
    summary.append([])
    summary.append(['Как работать', 'Сначала откройте лист «Требуют проверки». Строка там не означает ошибку — это сигнал для ручной сверки.', '', ''])
    summary.append(['Что проверяется', 'Скачок индивидуального расхода; уменьшение показания; дата вне срока работы счётчика; будущее число; несоответствие счётчика, даты, значения или статуса связанной заявки контролёра; несколько заявок на одно показание.', '', ''])
    summary.append(['Исправление', 'Используйте ссылку «Исправить» в строке или отдельную кнопку в карточке показания. Перед записью система покажет «было → станет».', '', ''])
    summary.append(['Проверка в браузере', 'Открыть страницу проверки показаний', '', ''])
    review_url = request.build_absolute_uri(reverse('water_readings_review'))
    summary['B12'].hyperlink = f'{review_url}?threshold={threshold_text}'
    summary['B12'].style = 'Hyperlink'
    summary.column_dimensions['A'].width = 46
    summary.column_dimensions['B'].width = 100
    for row in summary.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical='top', wrap_text=True)
        row[1].alignment = Alignment(vertical='top', wrap_text=True)

    review_sheet = wb.create_sheet('Требуют проверки')
    review_sheet.append(XLSX_HEADERS)
    for row in flagged:
        _append_xlsx_row(review_sheet, row, request)
    _style_data_sheet(review_sheet, 'ReadingsToReview')

    all_sheet = wb.create_sheet('Все показания')
    all_sheet.append(XLSX_HEADERS)
    for row in rows:
        _append_xlsx_row(all_sheet, row, request)
    _style_data_sheet(all_sheet, 'AllReadings')

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    LogEntry.objects.create(
        user=request.user,
        content_type=ContentType.objects.get_for_model(Reading),
        object_id='',
        object_repr='Выгрузка проверки показаний XLSX',
        action_flag=2,
        change_message=(
            f'Экспорт XLSX: порог {threshold_text} м³; '
            f'{len(rows)} записей; требуют проверки: {len(flagged)}'
        ),
    )

    response = HttpResponse(
        stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="trud-water-readings-review.xlsx"'
    response['Cache-Control'] = 'no-store'
    return response
