#!/usr/bin/env python3
"""Rebuild staged water readings from the legacy workbook without guessing dates."""

from __future__ import annotations

import argparse
from pathlib import Path
import re

from openpyxl import load_workbook


READING_COLUMNS = {
    'Основные счетчики': ((3, 'июнь, начальное'), (4, 'июнь, текущее'), (8, 'июль, начальное'), (9, 'июль, текущее'), (13, 'август, начальное'), (14, 'август, текущее')),
    'Полив': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Линия Миндальная (Миша) гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Миндальная Белорус гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Горная линия Харитоновой гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Дачная Усольцева гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Линия Останина гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Морская гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Садовая гидр ': ((4, 'июнь'), (5, 'июль'), (6, 'август')),
    'Лесная без гидр': ((3, 'июнь, предыдущее'), (4, 'июнь, текущее'), (8, 'июль, предыдущее'), (9, 'июль, текущее'), (13, 'август, предыдущее'), (14, 'август, текущее')),
    'линия Волкова без гидр': ((5, 'июнь, предыдущее'), (6, 'июнь, текущее'), (10, 'июль, предыдущее'), (11, 'июль, текущее'), (15, 'август, предыдущее'), (16, 'август, текущее')),
    'Линия Гузева без гидр': ((5, 'март'), (6, 'апрель'), (7, 'май'), (8, 'июнь (1)'), (9, 'июнь (2)'), (10, 'август')),
}

PHONE_RE = re.compile(r'(?<!\d)(?:\+?7|8)[\s()\-]*(?:\d[\s()\-]*){10}(?!\d)')


def table(sheet):
    headers = [cell.value for cell in sheet[1]]
    return headers, {name: number + 1 for number, name in enumerate(headers)}


def normalized_phone(value):
    digits = re.sub(r'\D', '', str(value))
    if len(digits) == 11 and digits[0] in '78':
        return '+7' + digits[1:]
    return str(value).strip()


def source_record(sheet, row_number):
    parts = []
    for column in range(2, min(sheet.max_column, 4) + 1):
        value = sheet.cell(row_number, column).value
        if value not in (None, ''):
            parts.append(str(value).strip())
    return ' | '.join(parts)


def append_note(existing, note):
    existing = str(existing or '').strip()
    if note in existing:
        return existing
    return f'{existing}\n{note}'.strip()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('package', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()

    source_formula = load_workbook(args.source, data_only=False, read_only=True)
    source_values = load_workbook(args.source, data_only=True, read_only=True)
    package = load_workbook(args.package)

    plot_sheet = package['Участки']
    _, plot_columns = table(plot_sheet)
    plots = {}
    for row_number in range(2, plot_sheet.max_row + 1):
        plot_id = plot_sheet.cell(row_number, plot_columns['plot_id']).value
        if plot_id:
            plots[plot_id] = row_number

    people_sheet = package['Люди']
    _, people_columns = table(people_sheet)
    people = {}
    for row_number in range(2, people_sheet.max_row + 1):
        plot_id = people_sheet.cell(row_number, people_columns['plot_id']).value
        if plot_id:
            people[plot_id] = row_number

    phone_updates = 0
    source_notes = 0
    for plot_id, plot_row in plots.items():
        sheet_name = plot_sheet.cell(plot_row, plot_columns['source_sheet']).value
        source_row = plot_sheet.cell(plot_row, plot_columns['source_row']).value
        if sheet_name not in source_formula.sheetnames or not source_row:
            continue
        record = source_record(source_formula[sheet_name], int(source_row))
        matches = PHONE_RE.findall(record)
        current_phone = plot_sheet.cell(plot_row, plot_columns['phone']).value
        if not current_phone and len(matches) == 1:
            current_phone = normalized_phone(matches[0])
            plot_sheet.cell(plot_row, plot_columns['phone'], current_phone)
            phone_updates += 1
        person_row = people.get(plot_id)
        if person_row:
            if current_phone and not people_sheet.cell(person_row, people_columns['phone']).value:
                people_sheet.cell(person_row, people_columns['phone'], current_phone)
            note = f'Исходная запись ({sheet_name}, строка {source_row}): {record}'
            cell = people_sheet.cell(person_row, people_columns['notes'])
            updated = append_note(cell.value, note)
            if updated != (cell.value or ''):
                cell.value = updated
                source_notes += 1

    meter_sources = []
    individual_sheet = package['Индивидуальные счетчики']
    _, individual_columns = table(individual_sheet)
    for row_number in range(2, individual_sheet.max_row + 1):
        meter_id = individual_sheet.cell(row_number, individual_columns['meter_id']).value
        plot_id = individual_sheet.cell(row_number, individual_columns['plot_id']).value
        plot_row = plots.get(plot_id)
        if meter_id and plot_row:
            meter_sources.append((
                meter_id,
                plot_sheet.cell(plot_row, plot_columns['source_sheet']).value,
                plot_sheet.cell(plot_row, plot_columns['source_row']).value,
            ))

    system_sheet = package['Общие и контрольные']
    _, system_columns = table(system_sheet)
    for row_number in range(2, system_sheet.max_row + 1):
        meter_sources.append((
            system_sheet.cell(row_number, system_columns['meter_id']).value,
            system_sheet.cell(row_number, system_columns['source_sheet']).value,
            system_sheet.cell(row_number, system_columns['source_row']).value,
        ))

    reading_sheet = package['Показания']
    if reading_sheet.max_row > 1:
        reading_sheet.delete_rows(2, reading_sheet.max_row - 1)

    readings = 0
    skipped_blanks = 0
    for meter_id, sheet_name, source_row in meter_sources:
        if not meter_id or sheet_name not in READING_COLUMNS or not source_row:
            continue
        sheet = source_values[sheet_name]
        for column, label in READING_COLUMNS[sheet_name]:
            value = sheet.cell(int(source_row), column).value
            if value in (None, ''):
                skipped_blanks += 1
                continue
            if not isinstance(value, (int, float)):
                continue
            reading_sheet.append([
                meter_id,
                '',
                value,
                f'Исходное состояние счётчика: {label}; точная дата неизвестна',
                sheet_name,
                int(source_row),
            ])
            readings += 1

    readme = package['README']
    readme.append(['Пересборка', 'Телефоны перенесены в phone; исходные записи сохранены в примечаниях; состояния счётчиков извлечены по структуре каждого листа без придуманных дат.'])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    package.save(args.output)
    source_formula.close()
    source_values.close()
    package.close()
    print(f'Готово: {args.output}')
    print(f'Телефонов дополнено: {phone_updates}; исходных записей сохранено: {source_notes}; показаний подготовлено: {readings}; пустых ячеек пропущено: {skipped_blanks}.')


if __name__ == '__main__':
    main()
